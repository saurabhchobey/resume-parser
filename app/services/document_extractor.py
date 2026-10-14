"""Safe text extraction for supported resume document formats."""

from io import BytesIO

import fitz
from docx import Document


class DocumentExtractionError(ValueError):
    """Raised when a supported file cannot yield usable text."""


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_text(filename: str, content: bytes) -> str:
    """Extract normalized text from a PDF or DOCX byte payload."""

    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix not in {"pdf", "docx"}:
        raise DocumentExtractionError("Only PDF and DOCX files are supported.")

    try:
        if suffix == "pdf":
            with fitz.open(stream=content, filetype="pdf") as document:
                text = "\n".join(page.get_text("text") for page in document)
        else:
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_rows = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
            text = "\n".join(paragraphs + table_rows)
    except Exception as exc:
        raise DocumentExtractionError("The uploaded document could not be read.") from exc

    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not normalized:
        raise DocumentExtractionError("The uploaded document contains no extractable text.")
    return normalized

